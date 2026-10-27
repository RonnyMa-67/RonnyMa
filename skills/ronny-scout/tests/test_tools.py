import json
import sys
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT/'scripts'))
from parse_leads import load_rows, canonicalize, split_markdown, META_EXPORT_FIELDS
from export_xlsx import export, priority, wa_link, HEADERS, row_values
from append_to_template import append
from validate_export import validate

class Tests(unittest.TestCase):
    def sample(self):
        return json.loads((ROOT/'examples/sample_research.json').read_text(encoding='utf-8'))

    def test_input_csv(self):
        records = canonicalize(load_rows(ROOT/'examples/sample_leads.csv'))
        self.assertEqual(len(records), 2)
        self.assertEqual(records[0]['lead_id'],'SYN-0001')
        self.assertFalse(records[0]['duplicate_exact_id'])

    def test_whatsapp(self):
        self.assertEqual(wa_link('+1 202 555 0127'), 'https://wa.me/12025550127')
        self.assertEqual(wa_link('https://wa.me/12025550127'), 'https://wa.me/12025550127')
        self.assertEqual(wa_link('12025550127'), '')
        self.assertEqual(wa_link('https://wa.me/1202555012712025550127'), '')

    def test_priority(self):
        self.assertEqual(priority('P1｜重点推荐'),'P1')
        self.assertEqual(priority('待核实'),'UNVERIFIED')
        with self.assertRaises(ValueError):priority('P9')

    def test_fixed_a_j(self):
        r=self.sample()
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'sample.xlsx'
            export(r,p)
            self.assertIn('10 columns A-J',validate(p))
            ws=load_workbook(p).active
            self.assertEqual([c.value for c in ws[1]],HEADERS['zh'])
            self.assertEqual(ws.max_column,10)
            self.assertEqual(ws['C3'].value,'示例国家')
            self.assertIn('主营：',ws['D2'].value)
            self.assertIn('背景：',ws['D2'].value)
            self.assertIn('电话：',ws['F3'].value)
            self.assertIn('邮箱：',ws['F3'].value)
            self.assertNotIn('官网：',ws['F2'].value)
            self.assertIn('官网：',ws['G2'].value)
            self.assertEqual(ws['F3'].hyperlink.target,'https://wa.me/12025550127')
            self.assertEqual(ws['G2'].hyperlink.target,'https://example.com/')
            self.assertEqual(ws['I2'].value,'P2｜值得沟通')
            self.assertTrue(ws['B2'].comment.text.startswith('原始线索ID'))
            self.assertTrue(ws['B3'].font.bold)
            self.assertTrue(ws['E3'].font.bold)

    def test_no_extra_link_columns(self):
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaises(ValueError):export(self.sample(),Path(d)/'bad.xlsx',link_mode='independent')

    def test_preserve_master(self):
        w=Workbook();s=w.active;s.title='研究主表'
        for j,h in enumerate(HEADERS['zh'],1):s.cell(1,j,h)
        for j in range(11,22):s.cell(1,j,f'Staff Field {j}')
        s['B2']='Preexisting Client';s['B2'].comment=Comment('原始线索ID：EXIST-001','Ops')
        s['K2']='DO_NOT_CHANGE';s['K2'].comment=Comment('Existing staff note','Ops')
        s['U2']='Supplier match'
        d=DataValidation(type='list',formula1='"yes,no"');s.add_data_validation(d);d.add('L2:L100')
        ws=w.create_sheet('看板');ws['A1']='A';ws['A2']=1;ws['A3']=2
        chart=BarChart();chart.add_data(Reference(ws,min_col=1,min_row=1,max_row=3));ws.add_chart(chart,'C2')
        w.create_sheet('历史客户');w.create_sheet('现场登记')
        with tempfile.TemporaryDirectory() as t:
            src=Path(t)/'master.xlsx';dst=Path(t)/'result.xlsx';w.save(src)
            append(src,self.sample(),dst)
            a=load_workbook(src);b=load_workbook(dst)
            self.assertEqual(a.sheetnames,b.sheetnames)
            self.assertEqual(len(a['看板']._charts),len(b['看板']._charts))
            self.assertEqual(len(a['研究主表'].data_validations.dataValidation),len(b['研究主表'].data_validations.dataValidation))
            self.assertEqual(b['研究主表']['K2'].value,'DO_NOT_CHANGE')
            self.assertEqual(b['研究主表']['K2'].comment.text,'Existing staff note')
            self.assertEqual(b['研究主表']['U2'].value,'Supplier match')
            self.assertIsNone(b['研究主表']['K3'].value)
            self.assertIsNone(b['研究主表']['U3'].value)
            self.assertEqual(b['研究主表']['D3'].value,row_values(self.sample()[0])[3])
            self.assertEqual(b['研究主表']['F4'].hyperlink.target,'https://wa.me/12025550127')
            with self.assertRaises(ValueError):append(dst,self.sample(),Path(t)/'dup.xlsx')

    def test_mismatched_master_fails(self):
        w=Workbook();w.active['A1']='Customer'
        with tempfile.TemporaryDirectory() as t:
            p=Path(t)/'wrong.xlsx';w.save(p)
            with self.assertRaises(ValueError):append(p,self.sample(),Path(t)/'out.xlsx')

    def test_duplicate_ids(self):
        r=self.sample();r[1]['lead_id']=r[0]['lead_id']
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaises(ValueError):export(r,Path(d)/'bad.xlsx')

    def test_md_pipes(self):
        self.assertEqual(split_markdown('| A | B\\|C |'),['A','B|C'])

    def test_headerless_meta_md_preserves_first(self):
        fields=['2036052327060828','10/26/26','ad-id','ad','adset-id','adset','campaign-id','campaign',
                'form-id','form','false','fb','importer_/_distributor','slush_machines','online_meeting','',
                'Sample Country','test@example.invalid','Example Name','Example Company','+1 202 555 0127','complete']
        self.assertEqual(len(fields),len(META_EXPORT_FIELDS))
        with tempfile.TemporaryDirectory() as td:
            md=Path(td)/'in.md';md.write_text('| '+' | '.join(fields)+' |\n| '+' | '.join(['---']*22)+' |\n')
            rows=canonicalize(load_rows(md))
            self.assertEqual(len(rows),1)
            self.assertEqual(rows[0]['lead_id'],'2036052327060828')
            self.assertEqual(rows[0]['product_interest'],'slush_machines')

if __name__=='__main__':unittest.main()