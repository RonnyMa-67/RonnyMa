#!/usr/bin/env python3
"""Validate the fixed A-J research export, hyperlinks and readable row layout."""
import argparse,re
from pathlib import Path
from urllib.parse import urlparse
from openpyxl import load_workbook
try:
    from export_xlsx import HEADERS
except ImportError:
    from scripts.export_xlsx import HEADERS


def validate(path,lang='zh'):
    errors=[]
    wb=load_workbook(path)
    if len(wb.sheetnames)!=1:errors.append('Standalone export must contain one sheet')
    ws=wb.active
    headers=[ws.cell(1,i).value for i in range(1,11)]
    if headers!=HEADERS[lang]:errors.append(f'Incorrect A-J order: {headers!r}')
    if ws.max_column!=10:errors.append(f'Expected exactly 10 visible columns, got {ws.max_column}')
    ids=set()
    for r in range(2,ws.max_row+1):
        if not ws.cell(r,2).value:errors.append(f'Row {r}: missing customer')
        c=str(ws.cell(r,3).value or '')
        if '\n' in c or re.search(r'区号|电话前缀|phone prefix',c,re.I):errors.append(f'Row {r}: country column contaminated')
        d=str(ws.cell(r,4).value or '')
        if d and ('\n' not in d or ('主营：' not in d and '背景：' not in d)):
            errors.append(f'Row {r}: company not segmented')
        for col in (4,8,10):
            v=str(ws.cell(r,col).value or '')
            if re.search(r'(?i)multiple_categories|need_selection_advice|hotel_/_|retail_/_',v):
                errors.append(f'Row {r}: raw form code leaked into prose {col}')
        p=str(ws.cell(r,9).value or '')
        if not (p.startswith(('P1','P2','P3')) or p in ('待核实','UNVERIFIED')):
            errors.append(f'Row {r}: invalid priority {p}')
        for col in range(1,11):
            cell=ws.cell(r,col)
            if cell.value and not cell.alignment.wrap_text:errors.append(f'{cell.coordinate}: wrapping disabled')
            if cell.hyperlink:
                target=cell.hyperlink.target or ''
                try:u=urlparse(target)
                except ValueError:u=None
                if not u or u.scheme not in {'http','https','mailto'}:errors.append(f'{cell.coordinate}: invalid hyperlink')
                if 'wa.me/' in target and not re.fullmatch(r'\d{7,15}',target.split('wa.me/',1)[1].split('?',1)[0]):
                    errors.append(f'{cell.coordinate}: invalid WA link')
        if ws.row_dimensions[r].height and ws.row_dimensions[r].height>190:errors.append(f'Row {r}: excessive height')
        cm=ws.cell(r,2).comment
        match=re.search(r'原始线索ID：([^\n]+)',cm.text) if cm else None
        if match:
            if match.group(1) in ids:errors.append(f'Row {r}: duplicate lead ID')
            ids.add(match.group(1))
    if not ws.auto_filter.ref:errors.append('Missing filter')
    if errors:raise ValueError('\n'.join(errors))
    return f'PASS: {ws.max_row-1} data rows, 10 columns A-J, {len(ids)} source IDs'

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--input',type=Path,required=True)
    p.add_argument('--lang',default='zh',choices=['zh','en']);a=p.parse_args();print(validate(a.input,a.lang))