#!/usr/bin/env python3
"""Create the fixed A-J B2B research table from already researched JSON records.

This tool formats supplied research. It does not perform web research.
"""
import argparse
import json
import re
from pathlib import Path
from urllib.parse import urlparse, parse_qs

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = {
    'zh': ['分配业务员', '客户信息', '国家', '公司介绍（主营·背景·规模）', '原始采购需求',
           '联系方式 / WhatsApp', '官网与社媒', '匹配与合作潜力', '背调评级', '开发切入点'],
    'en': ['Owner', 'Customer', 'Country', 'Company profile (business / background / scale)',
           'Original buying intent', 'Contacts / WhatsApp', 'Website & social',
           'Business fit & opportunity', 'Priority', 'Suggested next step'],
}
PRIORITY_MAP = {'P1': 'P1｜重点推荐', 'P2': 'P2｜值得沟通', 'P3': 'P3｜适合培育', 'UNVERIFIED': '待核实'}
COLORS = {'P1': ('DDF2EB', '157C60'), 'P2': ('E5EFFB', '24598E'),
          'P3': ('F0F0F0', '555555'), 'UNVERIFIED': ('FFF1DA', '996000')}
OWNER_COLORS = {'yellow':'F9E5A0','pink':'F5CAE0','blue':'C9E2F7',
                'green':'D5EBD9','orange':'F9D2B2'}
SOURCE_FIELDS = {2:'customer_name',3:'country',4:'company_profile',5:'original_purchase_intent',
                 6:'phone',7:'social_links',8:'cooperation_potential',9:'priority_level',
                 10:'development_suggestion'}


def valid_url(value):
    s = str(value or '').strip()
    if not s:
        return ''
    try:
        u = urlparse(s)
    except ValueError:
        return ''
    return s if u.scheme in {'http','https'} and u.netloc else ''


def wa_link(value):
    """A link is only generated for an explicit WhatsApp-labelled number or URL."""
    v = str(value or '').strip()
    if not v:
        return ''
    if 'wa.me/' in v:
        raw = v.split('wa.me/',1)[1].split('?',1)[0].split('/',1)[0]
    elif 'api.whatsapp.com/send?' in v:
        raw = parse_qs(urlparse(v).query).get('phone', [''])[0]
    elif v.startswith('+'):
        raw = v
    else:
        return ''
    digits = re.sub(r'\D','',raw)
    if not 7 <= len(digits) <= 15:
        return ''
    return 'https://wa.me/'+digits


def priority(value):
    v = str(value or '待核实').strip().upper()
    for rank in ('P1','P2','P3'):
        if re.match(r'^'+rank+r'(?=$|\W)', v):
            return rank
    if v in {'待核实','UNVERIFIED','UNKNOWN','PENDING','TO VERIFY'}:
        return 'UNVERIFIED'
    raise ValueError('Invalid priority: '+repr(value))


def display_country(record):
    """Keep only the country field; never derive it from phone code."""
    c = str(record.get('country') or '').strip()
    if '\n' in c or re.search(r'(区号|phone prefix|电话号|电话前缀)',c,re.I):
        raise ValueError('Country column should contain a country name only: '+repr(c))
    return c


def company_text(item):
    name=str(item.get('company_name') or '').strip()
    typ=str(item.get('business_type') or '').strip()
    intro='｜'.join(x for x in [name,typ] if x)
    lines=[intro] if intro else []
    for label,key in [('主营','company_profile'),('背景','business_background'),
                      ('规模','company_scale'),('渠道','company_channels')]:
        val=str(item.get(key) or '').strip()
        if val:
            val=re.sub(r'^(主营|背景|规模|渠道)[：:]\s*','',val)
            lines.append(label+'：'+val)
    return '\n'.join(lines)


def original_intent(item):
    chunks=[str(item.get('original_purchase_intent') or '').strip()]
    if item.get('original_quantity') and str(item['original_quantity']) not in chunks[0]:
        chunks.append('数量：'+str(item['original_quantity']))
    if item.get('purchase_time_window') and str(item['purchase_time_window']) not in ' '.join(chunks):
        chunks.append('周期：'+str(item['purchase_time_window']))
    if item.get('meeting_preference') and str(item['meeting_preference']) not in ' '.join(chunks):
        chunks.append('沟通：'+str(item['meeting_preference']))
    return '\n'.join(s for s in chunks if s)


def contacts_text(item):
    return '\n'.join(x for x in [
        '电话：'+str(item['phone']) if item.get('phone') else '',
        'WhatsApp：'+(wa_link(item.get('whatsapp')) or str(item['whatsapp'])) if item.get('whatsapp') else '',
        '邮箱：'+str(item['email']) if item.get('email') else '',
    ] if x)


def socials_text(item):
    groups=[]
    if item.get('website'):
        groups.append('官网：'+str(item['website']))
    socials=item.get('social_links') or {}
    if isinstance(socials,dict):
        groups.extend(str(k)+'：'+str(v) for k,v in socials.items() if v)
    return '\n'.join(groups)


def sources_comment(item,field):
    lines=[]
    if field=='customer_name' and item.get('lead_id'):
        lines.append('原始线索ID：'+str(item['lead_id']))
    notes=item.get('field_notes') or {}
    subfields={'company_profile':['company_name','business_type','business_background','company_scale','company_channels'],
               'phone':['whatsapp','email'],'social_links':['website']}
    for k in [field]+subfields.get(field,[]):
        if isinstance(notes,dict) and notes.get(k):
            lines.append(k+'：'+str(notes[k]))
    if field in {'company_profile','phone','social_links','cooperation_potential','priority_level','development_suggestion'}:
        for s in item.get('sources') or []:
            if isinstance(s,dict) and valid_url(s.get('url')):
                lines.append(str(s.get('claim') or '来源')+'：'+valid_url(s['url']))
    if field in {'customer_name','original_purchase_intent'} and item.get('raw'):
        lines.append('原始字段：'+json.dumps(item['raw'],ensure_ascii=False,default=str)[:12000])
    return '\n'.join(lines)[:30000]


def row_values(item,lang='zh'):
    code=priority(item.get('priority_level'))
    name='\n'.join(x for x in [str(item.get('customer_name') or '').strip(),
                              str(item.get('company_name') or '').strip()] if x)
    return [str(item.get('assignee') or ''),name,display_country(item),company_text(item),
            original_intent(item),contacts_text(item),socials_text(item),
            str(item.get('cooperation_potential') or ''),
            PRIORITY_MAP[code] if lang=='zh' else code,
            str(item.get('development_suggestion') or '')]


def fill_row(ws,row_index,item,lang='zh'):
    values=row_values(item,lang)
    for j,value in enumerate(values,1):
        cell=ws.cell(row_index,j,value)
        cell.alignment=Alignment(wrap_text=True,vertical='top',horizontal='left')
        cell.font=Font(name='Arial',size=10,color='263A4B')
        cell.fill=PatternFill('solid',fgColor='FFFFFF' if row_index%2==0 else 'F6F8FB')
        cell.border=Border(bottom=Side(style='hair',color='C6D7E3'))
        field=SOURCE_FIELDS.get(j)
        if field:
            note=sources_comment(item,field)
            if note:cell.comment=Comment(note,'Research notes')
    color=OWNER_COLORS.get(str(item.get('assignee_color') or '').lower())
    if color:ws.cell(row_index,1).fill=PatternFill('solid',fgColor=color)
    if item.get('onsite_meeting_interest'):
        for j in (2,5):
            ws.cell(row_index,j).font=Font(name='Arial',size=10,color='AD2534',bold=True)
    p=priority(item.get('priority_level'))
    fill,fg=COLORS[p]
    ws.cell(row_index,9).fill=PatternFill('solid',fgColor=fill)
    ws.cell(row_index,9).font=Font(name='Arial',size=10,color=fg,bold=True)
    # One hyperlink per cell. F contains all contact text and one optional WA link.
    direct=wa_link(item.get('whatsapp'))
    if direct:ws.cell(row_index,6).hyperlink=direct
    # G contains all website/social URLs and one main verified official website link if given.
    site=valid_url(item.get('website'))
    if site:ws.cell(row_index,7).hyperlink=site
    row_height=60 + min(108,max(len(values[3].splitlines()),len(values[5].splitlines()),len(values[6].splitlines()),len(values[7].splitlines()))*14)
    ws.row_dimensions[row_index].height=min(172,max(80,row_height))
    return values


def export(data,output,lang='zh',link_mode='compact'):
    if link_mode!='compact':
        raise ValueError('Fixed A-J layout disallows extra link columns; grouped F/G links are standard')
    if not isinstance(data,list) or not data:
        raise ValueError('Expected nonempty list of researched objects')
    ids=[str(x.get('lead_id')) for x in data if x.get('lead_id')]
    if len(ids)!=len(set(ids)):
        raise ValueError('Duplicate lead IDs')
    wb=Workbook();ws=wb.active;ws.title='客户背调' if lang=='zh' else 'Lead Research'
    ws.sheet_view.showGridLines=False;ws.sheet_view.zoomScale=85;ws.freeze_panes='D2'
    widths=[16,29,16,64,51,38,52,56,18,51]
    for j,head in enumerate(HEADERS[lang],1):
        cell=ws.cell(1,j,head)
        cell.fill=PatternFill('solid',fgColor='1F3B61')
        cell.font=Font(name='Arial',size=10,color='FFFFFF',bold=True)
        cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width=widths[j-1]
    ws.row_dimensions[1].height=34
    for i,item in enumerate(data,2):fill_row(ws,i,item,lang)
    ws.auto_filter.ref=f'A1:J{ws.max_row}'
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True);wb.save(output)
    return output


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--input',type=Path,required=True)
    p.add_argument('--output',type=Path,required=True)
    p.add_argument('--lang',choices=['zh','en'],default='zh')
    a=p.parse_args()
    rows=json.loads(a.input.read_text(encoding='utf-8'))
    out=export(rows,a.output,a.lang)
    print(f'Exported {len(rows)} records, fixed 10 columns A-J: {out}')

if __name__=='__main__':main()